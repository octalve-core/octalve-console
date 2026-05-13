# reports/report_generators.py

import csv
import io
import json
from datetime import datetime
from django.db.models import Sum, Count, Q
from django.utils import timezone
from django.core.files.storage import default_storage
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from .models import ReportInstance

class ReportGenerator:
    """Base class for report generation"""
    
    def __init__(self, report_instance):
        self.report = report_instance
        self.definition = report_instance.definition
        self.data = []
        self.headers = []
    
    def fetch_data(self):
        """Fetch data based on report type"""
        if self.definition.report_type == 'sales':
            self.fetch_sales_data()
        elif self.definition.report_type == 'revenue':
            self.fetch_revenue_data()
        elif self.definition.report_type == 'invoices':
            self.fetch_invoices_data()
        elif self.definition.report_type == 'payments':
            self.fetch_payments_data()
        elif self.definition.report_type == 'services':
            self.fetch_services_data()
        elif self.definition.report_type == 'customers':
            self.fetch_customers_data()
    
    def fetch_sales_data(self):
        """Fetch sales report data"""
        from orders.models import Order
        from billing.models import Invoice
        
        orders = Order.objects.filter(
            created_at__gte=self.report.period_start,
            created_at__lte=self.report.period_end,
            status='paid'
        )
        
        self.headers = ['Date', 'Order #', 'Customer', 'Product', 'Amount', 'Status']
        
        for order in orders:
            for item in order.items.all():
                self.data.append([
                    order.created_at.strftime('%Y-%m-%d'),
                    order.order_number,
                    order.client.email,
                    item.description,
                    float(item.unit_price),
                    order.status
                ])
    
    def fetch_revenue_data(self):
        """Fetch revenue report data"""
        from billing.models import Invoice, Transaction
        
        invoices = Invoice.objects.filter(
            paid_at__gte=self.report.period_start,
            paid_at__lte=self.report.period_end,
            status='paid'
        )
        
        self.headers = ['Date', 'Invoice #', 'Customer', 'Amount', 'Payment Method', 'Gateway']
        
        for invoice in invoices:
            transaction = invoice.transactions.first()
            self.data.append([
                invoice.paid_at.strftime('%Y-%m-%d'),
                invoice.invoice_number,
                invoice.client.email,
                float(invoice.total),
                transaction.payment_method if transaction else 'N/A',
                transaction.gateway.name if transaction and transaction.gateway else 'N/A'
            ])
    
    def fetch_invoices_data(self):
        """Fetch invoices report data"""
        from billing.models import Invoice
        
        invoices = Invoice.objects.filter(
            created_at__gte=self.report.period_start,
            created_at__lte=self.report.period_end
        )
        
        self.headers = ['Invoice #', 'Customer', 'Issue Date', 'Due Date', 'Status', 'Total', 'Paid', 'Due']
        
        for invoice in invoices:
            self.data.append([
                invoice.invoice_number,
                invoice.client.email,
                invoice.issue_date.strftime('%Y-%m-%d'),
                invoice.due_date.strftime('%Y-%m-%d'),
                invoice.status,
                float(invoice.total),
                float(invoice.amount_paid),
                float(invoice.amount_due)
            ])
    
    def fetch_payments_data(self):
        """Fetch payments report data"""
        from payments.models import Transaction
        
        transactions = Transaction.objects.filter(
            paid_at__gte=self.report.period_start,
            paid_at__lte=self.report.period_end,
            status='successful'
        )
        
        self.headers = ['Date', 'Reference', 'Customer', 'Amount', 'Gateway', 'Fee', 'Net']
        
        for transaction in transactions:
            self.data.append([
                transaction.paid_at.strftime('%Y-%m-%d %H:%M'),
                transaction.reference,
                transaction.client.email,
                float(transaction.amount),
                transaction.gateway.name if transaction.gateway else 'N/A',
                float(transaction.fee),
                float(transaction.net_amount)
            ])
    
    def fetch_services_data(self):
        """Fetch services report data"""
        from services.models import Service
        
        services = Service.objects.filter(
            created_at__gte=self.report.period_start,
            created_at__lte=self.report.period_end
        )
        
        self.headers = ['Service', 'Customer', 'Type', 'Status', 'Amount', 'Billing Cycle', 'Created']
        
        for service in services:
            self.data.append([
                service.label,
                service.client.email,
                service.service_type,
                service.status,
                float(service.amount),
                service.billing_cycle,
                service.created_at.strftime('%Y-%m-%d')
            ])
    
    def fetch_customers_data(self):
        """Fetch customers report data"""
        from django.contrib.auth import get_user_model
        User = get_user_model()
        
        users = User.objects.filter(
            date_joined__gte=self.report.period_start,
            date_joined__lte=self.report.period_end,
            role='client'
        )
        
        self.headers = ['Email', 'Full Name', 'Phone', 'Joined', 'Status', 'Total Spent']
        
        for user in users:
            from billing.models import Invoice
            total_spent = Invoice.objects.filter(client=user, status='paid').aggregate(Sum('total'))['total__sum'] or 0
            self.data.append([
                user.email,
                user.full_name,
                user.phone_number or 'N/A',
                user.date_joined.strftime('%Y-%m-%d'),
                'Active' if user.is_active else 'Inactive',
                float(total_spent)
            ])
    
    def generate_csv(self):
        """Generate CSV file"""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(self.headers)
        writer.writerows(self.data)
        
        # Save to storage
        file_name = f"reports/{self.report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        path = default_storage.save(file_name, io.BytesIO(output.getvalue().encode('utf-8')))
        
        return path, len(self.data)
    
    def generate_excel(self):
        """Generate Excel file"""
        wb = Workbook()
        ws = wb.active
        ws.title = self.definition.name[:31]
        
        # Add title
        title = ws.cell(row=1, column=1, value=self.report.title)
        title.font = Font(size=14, bold=True)
        
        # Add date range
        date_range = ws.cell(row=2, column=1, value=f"Period: {self.report.period_start.strftime('%Y-%m-%d')} to {self.report.period_end.strftime('%Y-%m-%d')}")
        date_range.font = Font(size=10)
        
        # Add headers
        for col, header in enumerate(self.headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
        
        # Add data
        for row, data_row in enumerate(self.data, 5):
            for col, value in enumerate(data_row, 1):
                ws.cell(row=row, column=col, value=value)
        
        # Adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save file
        file_name = f"reports/{self.report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output = io.BytesIO()
        wb.save(output)
        path = default_storage.save(file_name, output)
        
        return path, len(self.data)
    
    def generate_pdf(self):
        """Generate PDF file"""
        doc = SimpleDocTemplate(io.BytesIO(), pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        story = []
        
        # Title
        title = Paragraph(f"<b>{self.report.title}</b>", styles['Title'])
        story.append(title)
        story.append(Spacer(1, 12))
        
        # Date range
        date_text = f"Period: {self.report.period_start.strftime('%Y-%m-%d')} to {self.report.period_end.strftime('%Y-%m-%d')}"
        story.append(Paragraph(date_text, styles['Normal']))
        story.append(Spacer(1, 12))
        
        # Table
        table_data = [self.headers] + self.data[:1000]  # Limit to 1000 rows for PDF
        
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        
        story.append(table)
        
        # Generate PDF
        file_name = f"reports/{self.report.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(letter))
        doc.build(story)
        path = default_storage.save(file_name, output)
        
        return path, len(self.data)
    
    def generate(self):
        """Generate report in the specified format"""
        self.fetch_data()
        
        if self.report.file_format == 'csv':
            return self.generate_csv()
        elif self.report.file_format == 'excel':
            return self.generate_excel()
        elif self.report.file_format == 'pdf':
            return self.generate_pdf()
        else:
            return self.generate_csv()
